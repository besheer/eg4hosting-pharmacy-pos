"""Catalog exchange and branch-scoped operational reports, v0.6."""
import csv, io, json, hashlib, secrets, sqlite3, zipfile
import datetime as dt
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from html import escape as esc
from pathlib import Path
from urllib.parse import urlencode
from fastapi import Request, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import Response, RedirectResponse
from sqlalchemy import select, func, text
from sqlalchemy.orm import Session, Mapped, mapped_column
from sqlalchemy import String, Text, Integer, DateTime
from .main import (app, Base, engine, db, Product, Branch, Batch, Sale, SaleItem,
    SalesReturn, Purchase, StockMove, User, Customer, Supplier, Audit,
    current_user, allowed, page, PERM_PREFIX)
from .pos_v05 import SaleTotals, ItemAmounts, Payment, cents

COLUMNS=['barcode','trade_name','active_ingredient','strength','dosage_form','manufacturer','sale_price','reorder_level']
LABELS=['الباركود','اسم الصنف','المادة الفعالة','التركيز','الشكل الدوائي','الشركة','سعر البيع','حد الطلب']
ALIASES=dict(zip(LABELS,COLUMNS))
ALIASES.update({'اسم الدواء':'trade_name','السعر':'sale_price','barcode':'barcode','name':'trade_name','price':'sale_price'})
class ImportJob(Base):
    __tablename__='catalog_import_jobs'
    token:Mapped[str]=mapped_column(String(64),primary_key=True)
    user_id:Mapped[int]=mapped_column(Integer)
    created_at:Mapped[dt.datetime]=mapped_column(DateTime,default=dt.datetime.utcnow)
    status:Mapped[str]=mapped_column(String(20),default='preview')
    payload:Mapped[str]=mapped_column(Text)
Base.metadata.create_all(engine)
PERM_PREFIX.update({'/catalog':'users'})


def user(request,s,permission):
    u=current_user(request,s)
    if not u or not allowed(s,u,permission):raise HTTPException(403,'غير مسموح')
    return u

def table(headers,rows):
    return '<div class="table-scroll"><table><thead><tr>'+''.join('<th>'+esc(str(x))+'</th>' for x in headers)+'</tr></thead><tbody>'+(''.join('<tr>'+''.join('<td>'+esc('—' if v is None else str(v))+'</td>' for v in row)+'</tr>' for row in rows) or '<tr><td colspan="20">لا توجد بيانات لهذه الفلاتر</td></tr>')+'</tbody></table></div>'

def workbook_bytes(headers,rows):
    # Runtime Excel export for the installed Python application.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='Data';ws.sheet_view.rightToLeft=True
    ws.append(headers)
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163A4A');c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=32
    for col in range(1,len(headers)+1):
        ws.column_dimensions[get_column_letter(col)].width=min(42,max(18,len(str(headers[col-1]))+4))
    out=io.BytesIO();wb.save(out);return out.getvalue()

def export(headers,rows,fmt,name):
    if fmt=='xlsx':return Response(workbook_bytes(headers,rows),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{name}.xlsx"'})
    if fmt!='csv':raise HTTPException(400,'اختر csv أو xlsx')
    out=io.StringIO(newline='');writer=csv.writer(out);writer.writerow(headers)
    for row in rows:
        # CSV spreadsheet formula injection protection. XLSX uses literal text.
        writer.writerow([("'"+x) if isinstance(x,str) and x.lstrip().startswith(('=','+','-','@','\t','\r')) else x for x in row])
    return Response(out.getvalue().encode('utf-8-sig'),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':f'attachment; filename="{name}.csv"'})

def snapshot(p):return {k:getattr(p,k) for k in COLUMNS} if p else None

def parse_upload(data,name):
    if len(data)>5*1024*1024:raise HTTPException(400,'الحد الأقصى 5 MB')
    if name.lower().endswith('.csv'):
        try:raw=data.decode('utf-8-sig')
        except UnicodeDecodeError:raise HTTPException(400,'احفظ CSV بترميز UTF-8')
        try:delimiter=csv.Sniffer().sniff(raw[:8192],delimiters=',;\t').delimiter
        except csv.Error:delimiter=','
        rows=list(csv.reader(io.StringIO(raw),delimiter=delimiter))
    elif name.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(x.file_size for x in z.infolist())>25*1024*1024:raise ValueError('حجم الملف بعد الفك كبير')
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            ws=wb.active
            if ws.max_row>10001 or ws.max_column>30:raise ValueError('الحد 10000 صنف و30 عمود')
            rows=[]
            for cells in ws.iter_rows():
                if any(c.data_type=='f' for c in cells):raise ValueError('الصيغ غير مسموحة؛ الصق القيم فقط')
                rows.append([c.value for c in cells])
            wb.close()
        except Exception as e:raise HTTPException(400,'تعذر قراءة Excel: '+str(e))
    else:raise HTTPException(400,'الصيغ المدعومة CSV وXLSX؛ احفظ ملفات XLS القديمة بصيغة XLSX')
    if not rows or len(rows)>10001:raise HTTPException(400,'الملف فارغ أو أكثر من 10000 صنف')
    headers=[ALIASES.get(str(x or '').strip(),str(x or '').strip().lower()) for x in rows[0]]
    if len(set(headers))!=len(headers) or any(x not in COLUMNS for x in headers):raise HTTPException(400,'أسماء أعمدة مكررة أو غير معروفة؛ استخدم القالب')
    if not {'barcode','trade_name','sale_price'}.issubset(headers):raise HTTPException(400,'الباركود والاسم وسعر البيع مطلوبة')
    return headers,rows[1:]

@app.get('/catalog')
def catalog(request:Request,q:str='',s:Session=Depends(db)):
    user(request,s,'users')
    stmt=select(Product).order_by(Product.trade_name)
    if q:stmt=stmt.where(Product.trade_name.contains(q)|Product.barcode.contains(q)|Product.active_ingredient.contains(q))
    rows=s.scalars(stmt.limit(200)).all()
    links=' '.join(f'<a class="action" href="/catalog/export?{urlencode({"fmt":fmt,"q":q})}">تصدير {fmt.upper()}</a>' for fmt in ['csv','xlsx'])
    body=f'''<div class="card"><h2>استيراد وتصدير الأصناف</h2><p>ارفع الملف لمعاينته أولًا. المطابقة بالباركود؛ لا يتغير رصيد المخزون من استيراد الأصناف.</p>
    <p><a href="/catalog/template?fmt=xlsx">قالب Excel</a> · <a href="/catalog/template?fmt=csv">قالب CSV</a> · {links}</p>
    <form method="post" action="/catalog/preview" enctype="multipart/form-data"><input type="file" name="file" accept=".csv,.xlsx" required><select name="mode"><option value="add">إضافة جديدة فقط؛ رفض الموجود</option><option value="upsert">إضافة وتحديث الموجود بالباركود</option></select><button>معاينة الاستيراد</button></form>
    <p class="muted">الحد 10000 صنف / 5 MB. احفظ الباركود كنص في Excel لحماية الأصفار الأولى. الأعمدة الاختيارية الغائبة لا تغيّر الموجود.</p></div>
    <form><input name="q" value="{esc(q,quote=True)}" placeholder="بحث بالاسم أو الباركود أو المادة"><button>بحث</button></form><p>عرض أول 200 نتيجة؛ التصدير يشمل كل نتائج البحث.</p>'''
    return page('إدارة الأصناف والملفات',body+table(LABELS,[[getattr(p,k) for k in COLUMNS] for p in rows]))

@app.get('/catalog/template')
def catalog_template(request:Request,fmt:str='xlsx',s:Session=Depends(db)):
    user(request,s,'users')
    return export(COLUMNS,[['0001234567890','صنف تجريبي — احذف هذا الصف','مادة تجريبية','500 mg','Tablets','Demo',25.5,5]],fmt,'products-template')

@app.get('/catalog/export')
def catalog_export(request:Request,fmt:str='xlsx',q:str='',s:Session=Depends(db)):
    user(request,s,'users');stmt=select(Product).order_by(Product.id)
    if q:stmt=stmt.where(Product.trade_name.contains(q)|Product.barcode.contains(q)|Product.active_ingredient.contains(q))
    return export(COLUMNS,[[getattr(p,k) for k in COLUMNS] for p in s.scalars(stmt)],fmt,'products')

@app.post('/catalog/preview')
async def preview(request:Request,file:UploadFile=File(...),mode:str=Form('add'),s:Session=Depends(db)):
    u=user(request,s,'users')
    if mode not in ('add','upsert'):raise HTTPException(400,'وضع استيراد غير صحيح')
    data=await file.read(5*1024*1024+1);headers,raw=parse_upload(data,file.filename or '')
    errors=[];planned=[];seen=set();existing={p.barcode:p for p in s.scalars(select(Product))}
    for n,values in enumerate(raw,2):
        if not any(v is not None and str(v).strip() for v in values):continue
        try:
            if len(values)>len(headers):raise ValueError('أعمدة زائدة')
            row=dict(zip(headers,values+[None]*(len(headers)-len(values))))
            if not isinstance(row['barcode'],str):raise ValueError('الباركود يجب أن يكون نصًا؛ اضبط العمود Text ثم أعد إدخاله')
            row={k:str(v).strip() if v is not None else '' for k,v in row.items()}
            if not row['barcode'] or not row['trade_name']:raise ValueError('الباركود والاسم مطلوبان')
            if row['barcode'] in seen:raise ValueError('باركود مكرر داخل الملف')
            seen.add(row['barcode'])
            limits={'barcode':64,'trade_name':160,'active_ingredient':160,'strength':80,'dosage_form':80,'manufacturer':120}
            for k,length in limits.items():
                if len(row.get(k,''))>length:raise ValueError(f'{k}: الحد {length} حرفًا')
                if any(ord(c)<32 for c in row.get(k,'')):raise ValueError('محارف تحكم غير مسموحة')
                if '<' in row.get(k,'') or '>' in row.get(k,''):raise ValueError('وسوم HTML غير مسموحة في بيانات الأصناف')
            price=Decimal(row['sale_price'])
            if not price.is_finite() or price<0 or price>100000000 or price.as_tuple().exponent < -2:raise ValueError('سعر غير صحيح؛ منزلتان عشريتان كحد أقصى')
            row['sale_price']=float(price)
            if 'reorder_level' in row:
                reorder=Decimal(row['reorder_level'] or '0')
                if not reorder.is_finite() or reorder!=reorder.to_integral_value() or not 0<=reorder<=1000000:raise ValueError('حد الطلب يجب أن يكون عددًا صحيحًا موجبًا أو صفرًا')
                row['reorder_level']=int(reorder)
            old=existing.get(row['barcode'])
            if old and mode=='add':raise ValueError('الباركود موجود؛ اختر وضع التحديث لتعديله')
            planned.append({'line':n,'values':row,'before':snapshot(old),'action':'تحديث' if old else 'إضافة'})
        except (ValueError,InvalidOperation) as e:errors.append([n,str(e)])
    if not planned and not errors:errors.append([0,'لا توجد أصناف'])
    payload={'rows':planned,'errors':errors,'mode':mode}
    job=ImportJob(token=secrets.token_urlsafe(32),user_id=u.id,payload=json.dumps(payload,ensure_ascii=False));s.add(job);s.commit()
    return RedirectResponse('/catalog/preview/'+job.token,303)

def job_for(token,request,s):
    u=user(request,s,'users');job=s.get(ImportJob,token)
    if not job or job.user_id!=u.id:raise HTTPException(404,'المعاينة غير موجودة')
    if dt.datetime.utcnow()-job.created_at>dt.timedelta(hours=24):raise HTTPException(410,'انتهت المعاينة؛ ارفع الملف مجددًا')
    return job,json.loads(job.payload)

@app.get('/catalog/preview/{token}')
def preview_page(token:str,request:Request,s:Session=Depends(db)):
    job,data=job_for(token,request,s);rows=data['rows'];errors=data['errors']
    body=f'<div class="card"><p>صالح: {len(rows)} · أخطاء: {len(errors)} · الحالة: {esc(job.status)}</p>'
    if errors:body+='<p class="warn">لم يتم استيراد أي بيانات. أصلح جميع الأخطاء وأعد رفع الملف.</p>'+f'<a href="/catalog/errors/{token}">تنزيل الأخطاء CSV</a>'+table(['صف الملف','الخطأ'],errors[:200])
    elif job.status=='preview':body+=f'<p>راجع المعاينة ثم اعتمد العملية. سيتم حفظ نسخة احتياطية قبل التغيير.</p><form method="post" action="/catalog/confirm/{token}"><button>اعتماد استيراد {len(rows)} صنف</button></form>'
    body+='<a href="/catalog">العودة للأصناف</a></div><h2>معاينة أول 200 صنف</h2>'
    return page('معاينة الاستيراد',body+table(['الصف','الإجراء']+LABELS,[[r['line'],r['action']]+[r['values'].get(k,'بدون تغيير' if r['before'] else '') for k in COLUMNS] for r in rows[:200]]))

@app.get('/catalog/errors/{token}')
def import_errors(token:str,request:Request,s:Session=Depends(db)):
    _,data=job_for(token,request,s);return export(['row','error'],data['errors'],'csv','import-errors')

@app.post('/catalog/confirm/{token}')
def confirm(token:str,request:Request,s:Session=Depends(db)):
    if engine.dialect.name!='sqlite':raise HTTPException(400,'الاستيراد المعتمد في هذا الإصدار يدعم SQLite فقط')
    # Read first, back up from a separate SQLite connection, then recheck under writer lock.
    job,data=job_for(token,request,s)
    if job.status=='done':return RedirectResponse('/catalog/preview/'+token,303)
    if data['errors']:raise HTTPException(400,'أصلح أخطاء الملف أولًا')
    s.rollback()
    out=Path('backups');out.mkdir(exist_ok=True)
    backup=out/f'pre-import-{dt.datetime.now():%Y%m%d-%H%M%S}-{secrets.token_hex(4)}.db'
    with sqlite3.connect(engine.url.database) as source,sqlite3.connect(backup) as target:source.backup(target)
    s.execute(text('BEGIN IMMEDIATE'));job,data=job_for(token,request,s)
    if job.status=='done':return RedirectResponse('/catalog/preview/'+token,303)
    for row in data['rows']:
        p=s.scalar(select(Product).where(Product.barcode==row['values']['barcode']))
        if snapshot(p)!=row['before']:raise HTTPException(409,'تغيرت بيانات الأصناف منذ المعاينة؛ أعد رفع الملف')
        if p:
            for k,v in row['values'].items():setattr(p,k,v)
        else:s.add(Product(**row['values']))
    job.status='done';s.add(Audit(action='CATALOG_IMPORT',details=json.dumps({'count':len(data['rows']),'token':token,'backup':str(backup)},ensure_ascii=False)));s.commit()
    return RedirectResponse('/catalog/preview/'+token,303)

TITLES={'sales':'المبيعات والمرتجعات','profit':'مجمل الربح حسب الصنف','payments':'تحصيلات طرق الدفع','inventory':'المخزون والتشغيلات','low':'نواقص وإعادة الطلب','expiry':'قرب انتهاء الصلاحية','ledger':'حركة الأصناف','purchases':'فواتير المشتريات'}

def scope(request,s,branch_id):
    u=user(request,s,'reports')
    if not allowed(s,u,'users'):
        if not u.branch_id or branch_id and branch_id!=u.branch_id:raise HTTPException(403,'الفرع غير مسموح')
        branch_id=u.branch_id
    if branch_id and not s.get(Branch,branch_id):raise HTTPException(400,'الفرع غير موجود')
    return u,branch_id

def report_data(s,kind,start,end,branch_id,q):
    if start>end:raise HTTPException(400,'تاريخ البداية بعد النهاية')
    if kind not in TITLES:raise HTTPException(400,'نوع تقرير غير صحيح')
    branches={b.id:b.name for b in s.scalars(select(Branch))}
    products={p.id:p for p in s.scalars(select(Product))}
    def match(pid):
        p=products[pid];return not q or q.casefold() in ' '.join([p.trade_name,p.barcode,p.active_ingredient]).casefold()
    def within(date):return start<=date.date()<=end
    money=lambda n:round(n/100,2)
    if kind in ('sales','profit'):
        stmt=select(SaleItem,Sale).join(Sale,Sale.id==SaleItem.sale_id)
        if branch_id:stmt=stmt.where(Sale.branch_id==branch_id)
        amounts={a.item_id:a for a in s.scalars(select(ItemAmounts))}
        totals={t.sale_id:t for t in s.scalars(select(SaleTotals))}
        users={u.id:u.display_name or u.username for u in s.scalars(select(User))}
        returns={}
        for r in s.scalars(select(SalesReturn).order_by(SalesReturn.created_at,SalesReturn.id)):returns.setdefault(r.sale_item_id,[]).append(r)
        rows=[];groups={};net_sum=tax_sum=cost_sum=0;missing=0
        for item,sale in s.execute(stmt):
            if not match(item.product_id):continue
            p=products[item.product_id];a=amounts.get(item.id);total=totals.get(sale.id)
            events=[]
            if within(sale.created_at):events.append((sale.created_at,'بيع',item.qty,a.net if a else cents(item.unit_price)*item.qty,a.tax if a else 0,a.cost if a else None))
            cumulative=0
            for r in returns.get(item.id,[]):
                prev=cumulative;cumulative+=r.qty
                def portion(v):return int((Decimal(v)*cumulative/item.qty).quantize(Decimal(1),rounding=ROUND_HALF_UP))-int((Decimal(v)*prev/item.qty).quantize(Decimal(1),rounding=ROUND_HALF_UP))
                if within(r.created_at):events.append((r.created_at,'مرتجع',-r.qty,-portion(a.net) if a else -cents(r.amount),-portion(a.tax) if a else 0,-portion(a.cost) if a else None))
            for date,typ,qty,net,tax,cost in events:
                profit=net-cost if cost is not None else None
                net_sum+=net;tax_sum+=tax
                if cost is None:missing+=1
                else:cost_sum+=cost
                rows.append([date.strftime('%Y-%m-%d %H:%M'),sale.id,typ,branches[sale.branch_id],users.get(total.user_id,'غير مسجل') if total else 'غير مسجل',p.barcode,p.trade_name,qty,money(net),money(tax),money(net+tax),money(cost) if cost is not None else None,money(profit) if profit is not None else None])
                g=groups.setdefault(p.id,[p.barcode,p.trade_name,0,0,0,False]);g[2]+=qty;g[3]+=net
                if cost is None:g[5]=True
                else:g[4]+=cost
        summary={'صافي البيع بدون ضريبة':money(net_sum),'صافي الضريبة':money(tax_sum),'الإجمالي شامل الضريبة':money(net_sum+tax_sum),'مجمل الربح':None if missing else money(net_sum-cost_sum)}
        note='المرتجعات تُخصم في تاريخ إرجاعها. مجمل الربح قبل مصروفات التشغيل، بدون الضريبة، من تكلفة البيع المحفوظة. الفواتير القديمة بلا تكلفة محفوظة تظهر «—» ولا يُفترض لها ربح.'
        if kind=='profit':return ['الباركود','الصنف','صافي الكمية','صافي البيع بدون ضريبة','التكلفة','مجمل الربح'],[[g[0],g[1],g[2],money(g[3]),None if g[5] else money(g[4]),None if g[5] else money(g[3]-g[4])] for g in groups.values()],summary,note
        return ['التاريخ UTC','الفاتورة','النوع','الفرع','الكاشير','الباركود','الصنف','الكمية','بدون ضريبة','الضريبة','الإجمالي','التكلفة','مجمل الربح'],sorted(rows,key=lambda r:(r[0],r[1])),summary,note
    if kind=='payments':
        stmt=select(Payment,Sale).join(Sale,Sale.id==Payment.sale_id)
        if branch_id:stmt=stmt.where(Sale.branch_id==branch_id)
        rows=[[sale.created_at.strftime('%Y-%m-%d %H:%M'),sale.id,branches[sale.branch_id],p.method,money(p.amount)] for p,sale in s.execute(stmt) if within(sale.created_at)]
        return ['التاريخ UTC','الفاتورة','الفرع','طريقة الدفع','المبلغ'],rows,{'إجمالي التحصيل المسجل':round(sum(r[-1] for r in rows),2)},'تحصيلات v0.5 فأحدث فقط؛ لا تخصم التزامات المرتجعات لأنها لم تُصرف من بوابة دفع. بحث الصنف لا يطبق على هذا التقرير.'
    if kind in ('inventory','expiry','low'):
        stmt=select(Batch)
        if branch_id:stmt=stmt.where(Batch.branch_id==branch_id)
        batches=[b for b in s.scalars(stmt) if match(b.product_id)]
        if kind=='low':
            totals={}
            for b in batches:
                if b.expiry>dt.date.today():totals[(b.branch_id,b.product_id)]=totals.get((b.branch_id,b.product_id),0)+b.qty
            rows=[]
            for bid in ([branch_id] if branch_id else branches):
                for pid,p in products.items():
                    qty=totals.get((bid,pid),0)
                    if match(pid) and qty<=p.reorder_level:rows.append([branches[bid],p.barcode,p.trade_name,qty,p.reorder_level,max(0,p.reorder_level-qty)])
            return ['الفرع','الباركود','الصنف','رصيد صالح','حد الطلب','كمية للوصول للحد'],rows,{'أصناف تحتاج مراجعة':len(rows)},'رصيد لحظي صالح للبيع؛ يشمل الأصناف بلا تشغيلات. فلاتر التاريخ لا تطبق.'
        if kind=='expiry':batches=[b for b in batches if b.qty>0 and b.expiry<=dt.date.today()+dt.timedelta(days=90)]
        rows=[[branches[b.branch_id],products[b.product_id].barcode,products[b.product_id].trade_name,b.batch_no,b.expiry.isoformat(),b.qty,b.cost,round(b.qty*b.cost,2),'منتهي/غير صالح اليوم' if b.expiry<=dt.date.today() else 'صالح'] for b in sorted(batches,key=lambda b:b.expiry)]
        return ['الفرع','الباركود','الصنف','التشغيلة','الصلاحية','الكمية','تكلفة الوحدة','قيمة المخزون','الحالة'],rows,{'عدد التشغيلات':len(rows),'الوحدات':sum(r[5] for r in rows),'قيمة بالتكلفة':round(sum(r[7] for r in rows),2)},'رصيد لحظي، وليس رصيدًا تاريخيًا؛ فلاتر التاريخ لا تطبق. تقرير الصلاحية يشمل المنتهي والقريب خلال 90 يومًا.'
    if kind=='ledger':
        stmt=select(StockMove).order_by(StockMove.id)
        if branch_id:stmt=stmt.where(StockMove.branch_id==branch_id)
        rows=[[m.created_at.strftime('%Y-%m-%d %H:%M'),branches[m.branch_id],products[m.product_id].barcode,products[m.product_id].trade_name,m.batch_id,m.qty_delta,m.reason,m.ref] for m in s.scalars(stmt) if within(m.created_at) and match(m.product_id)]
        return ['التاريخ UTC','الفرع','الباركود','الصنف','تشغيلة ID','تغير الكمية','السبب','المرجع'],rows,{'عدد الحركات':len(rows),'صافي تغير الوحدات':sum(r[5] for r in rows)},'حركات الفترة فقط؛ ليست رصيد افتتاح وإقفال.'
    suppliers={p.id:p.name for p in s.scalars(select(Supplier))};stmt=select(Purchase).order_by(Purchase.id)
    if branch_id:stmt=stmt.where(Purchase.branch_id==branch_id)
    rows=[[p.created_at.strftime('%Y-%m-%d %H:%M'),p.id,p.invoice_no,branches[p.branch_id],suppliers.get(p.supplier_id,''),p.total] for p in s.scalars(stmt) if within(p.created_at)]
    return ['التاريخ UTC','ID','رقم الفاتورة','الفرع','المورد','الإجمالي'],rows,{'إجمالي فواتير المشتريات':round(sum(r[-1] for r in rows),2)},'قيمة الفواتير الأصلية قبل مرتجعات المورد. بحث الصنف لا يطبق.'

@app.get('/reports')
def reports(request:Request,kind:str='sales',start:dt.date|None=None,end:dt.date|None=None,branch_id:int|None=None,q:str='',fmt:str|None=None,s:Session=Depends(db)):
    u,branch_id=scope(request,s,branch_id);end=end or dt.date.today();start=start or end.replace(day=1)
    headers,rows,summary,note=report_data(s,kind,start,end,branch_id,q)
    if fmt=='print':
        return page(TITLES[kind],f'<p>الفترة: {start} إلى {end} · الفرع: {branch_id or "الكل"}</p><p>{esc(note)}</p><button onclick="print()">طباعة / حفظ PDF</button>'+table(headers,rows))
    if fmt:return export(headers,rows,fmt,'report-'+kind)
    opts=''.join(f'<option value="{k}" {"selected" if kind==k else ""}>{v}</option>' for k,v in TITLES.items())
    bo='<option value="">كل الفروع</option>' if allowed(s,u,'users') else ''
    bo+=''.join(f'<option value="{b.id}" {"selected" if b.id==branch_id else ""}>{esc(b.name)}</option>' for b in s.scalars(select(Branch)) if allowed(s,u,'users') or b.id==u.branch_id)
    # Empty branch selector is omitted by browser script so FastAPI receives None.
    filters=f'''<form class="card" method="get" onsubmit="if(!this.branch_id.value)this.branch_id.disabled=true"><label>التقرير <select name="kind">{opts}</select></label><label>من <input type="date" name="start" value="{start}"></label><label>إلى <input type="date" name="end" value="{end}"></label><label>الفرع <select name="branch_id">{bo}</select></label><input name="q" value="{esc(q,quote=True)}" placeholder="بحث صنف/باركود"><button>عرض التقرير</button></form>'''
    params={'kind':kind,'start':str(start),'end':str(end),'q':q}
    if branch_id:params['branch_id']=branch_id
    links=' '.join(f'<a class="action" href="/reports?{urlencode({**params,"fmt":f})}">تصدير {f.upper()}</a>' for f in ('csv','xlsx'))
    cards='<div class="grid">'+''.join(f'<div class="card"><div class="muted">{esc(k)}</div><div class="kpi">{v if v is not None else "غير مكتمل"}</div></div>' for k,v in summary.items())+'</div>'
    return page('مركز التقارير',filters+cards+f'<p>{esc(note)}</p><p class="no-print">{links} <a class="action" target="_blank" href="/reports?{urlencode({**params,"fmt":"print"})}">طباعة / PDF</a></p><p>عدد النتائج: {len(rows)}. الشاشة تعرض أول 500؛ التصدير يشمل كل النتائج المطابقة.</p>'+table(headers,rows[:500]))

@app.get('/')
def dashboard(request:Request,s:Session=Depends(db)):
    from .ui import icon
    u=current_user(request,s)
    bid=None if allowed(s,u,'users') else u.branch_id
    today=dt.date.today()
    def branch_filter(stmt,model):
        return stmt if allowed(s,u,'users') else stmt.where(model.branch_id==bid)
    sales=s.scalars(branch_filter(select(Sale).where(Sale.created_at>=dt.datetime.combine(today,dt.time.min)),Sale)).all()
    refunds=s.scalars(branch_filter(select(SalesReturn).where(SalesReturn.created_at>=dt.datetime.combine(today,dt.time.min)),SalesReturn)).all()
    exp=s.scalar(branch_filter(select(func.count()).select_from(Batch).where(Batch.qty>0,Batch.expiry<=today+dt.timedelta(days=90)),Batch))
    metrics=[('فواتير اليوم',len(sales),'عدد عمليات البيع'),('مبيعات اليوم',f'{sum(x.total for x in sales):,.2f}','شامل الضريبة · قبل المرتجعات'),('مرتجعات اليوم',f'{sum(x.amount for x in refunds):,.2f}','قيمة المرتجعات المسجلة'),('مراجعة الصلاحية',exp,'تشغيلات منتهية أو خلال 90 يومًا')]
    cards='<div class="grid">'+''.join(f'<div class="card metric"><div class="metric-label">{label}</div><div class="kpi">{value}</div><small>{note}</small></div>' for label,value,note in metrics)+'</div>'
    tiles=[('نقطة البيع','بيع جديد وإدارة السلة','/pos','pos','cart'),('استيراد وتصدير الأصناف','ملفات CSV وExcel','/catalog','users','file'),('مركز التقارير','المبيعات والربح والمخزون','/reports','reports','chart'),('المخزون والتشغيلات','الأرصدة والتكلفة والصلاحية','/inventory','inventory','box'),('العملاء والمرضى','بيانات العملاء والروشتات','/customers','pos','users'),('العروض والخصومات','إدارة عروض الكمية','/offers','users','tag')]
    modules='<div class="module-grid">'+''.join(f'<a class="card module" href="{url}"><span class="module-icon">{icon(symbol)}</span><div><h2>{title}</h2><small>{desc}</small></div><span class="module-arrow">←</span></a>' for title,desc,url,perm,symbol in tiles if allowed(s,u,perm))+'</div>'
    recent=s.scalars(branch_filter(select(Sale).order_by(Sale.id.desc()).limit(5),Sale)).all()
    recent_rows=[]
    for x in recent:
        recent_rows.append([f'#{x.id}',x.created_at.strftime('%Y-%m-%d %H:%M'),f'{x.total:,.2f}',{'cash':'نقدي','card':'بطاقة','wallet':'محفظة','bank':'تحويل','split':'دفع مقسّم'}.get(x.payment,x.payment)])
    latest=table(['الفاتورة','التاريخ UTC','الإجمالي','الدفع'],recent_rows) if recent_rows else '<div class="empty-state">لا توجد فواتير بعد.<br>ابدأ أول عملية من نقطة البيع.</div>'
    steps=''.join(f'<div class="step-row"><span class="step-number">{n}</span><div><b>{title}</b><p>{desc}</p></div></div>' for n,title,desc in [('01','جهّز دليل الأصناف','أضف الأدوية أو استورد ملف Excel.'),('02','سجّل المخزون','أدخل التشغيلات والصلاحية من المشتريات.'),('03','ابدأ البيع والمتابعة','استخدم نقطة البيع ثم راجع التقارير.')])
    action='<a class="action primary" href="/pos">+ عملية بيع جديدة</a>' if allowed(s,u,'pos') else ''
    scope_label='جميع الفروع' if allowed(s,u,'users') else 'الفرع المخصص للمستخدم'
    head=f'<div class="dashboard-head"><div><p>أهلًا، {esc(u.display_name or u.username)}</p><small class="muted">ملخص اليوم بتوقيت UTC · {scope_label}</small></div>{action}</div>'
    return page('لوحة التحكم',head+cards+'<div class="section-title"><h2>الوصول السريع</h2><span>وحدات العمل الأساسية</span></div>'+modules+f'<div class="dashboard-lower"><section class="card"><h2>آخر الفواتير</h2>{latest}</section><section class="card"><h2>خطوات البداية</h2>{steps}</section></div>')
