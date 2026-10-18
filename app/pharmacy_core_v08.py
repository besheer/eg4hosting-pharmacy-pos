"""Pharmacy Core v0.8: bulk Egyptian catalog import, smart equivalents and interaction checker."""
import csv, io, re, hashlib
from decimal import Decimal, InvalidOperation
from html import escape
from fastapi import Depends, Form, HTTPException, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse
from sqlalchemy import String, Integer, Text, select, or_
from sqlalchemy.orm import Mapped, mapped_column, Session
from openpyxl import load_workbook
from .main import app, Base, engine, db, Product, Audit, page, PERM_PREFIX
from .pharmacy_core_v07 import DrugProfile, DispensingPolicy, LEVELS

class DrugInteraction(Base):
    __tablename__='drug_interactions'
    id:Mapped[int]=mapped_column(primary_key=True)
    ingredient_a:Mapped[str]=mapped_column(String(180),index=True)
    ingredient_b:Mapped[str]=mapped_column(String(180),index=True)
    severity:Mapped[str]=mapped_column(String(20),default='moderate')
    note:Mapped[str]=mapped_column(Text,default='')
    source:Mapped[str]=mapped_column(String(240),default='')

Base.metadata.create_all(engine)
PERM_PREFIX.update({'/drug-import':'inventory','/smart-alternatives':'inventory','/interaction-checker':'pos','/drug-interactions':'users'})
SEVERITY={'minor':'🟡 بسيط','moderate':'🟠 متوسط','major':'🔴 شديد','contraindicated':'⛔ ممنوع معًا'}

def norm(v):
    return re.sub(r'\s+',' ',(v or '').strip()).casefold()

def cell(row,*names):
    low={norm(k):v for k,v in row.items() if k is not None}
    for n in names:
        v=low.get(norm(n))
        if v is not None and str(v).strip(): return str(v).strip()
    return ''

def rows_from_upload(name,data):
    if name.lower().endswith('.csv'):
        text=data.decode('utf-8-sig',errors='replace'); return list(csv.DictReader(io.StringIO(text)))
    if name.lower().endswith(('.xlsx','.xlsm')):
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True); ws=wb.active
        it=ws.iter_rows(values_only=True); headers=[str(x or '').strip() for x in next(it,())]
        return [dict(zip(headers,r)) for r in it]
    raise HTTPException(400,'الملف يجب أن يكون CSV أو XLSX')

def import_catalog(s,rows):
    made=updated=skipped=0
    for r in rows:
        barcode=cell(r,'barcode','gtin','باركود')
        trade=cell(r,'trade_name','english_name','commercial_name_en','name','اسم تجاري','الاسم التجاري')
        ar=cell(r,'arabic_name','commercial_name_ar','name_ar','الاسم العربي')
        generic=cell(r,'generic_name','active_ingredient','scientific_name','المادة الفعالة','الاسم العلمي')
        if not trade and not ar: skipped+=1; continue
        p=s.scalar(select(Product).where(Product.barcode==barcode)) if barcode else None
        if not p:
            p=s.scalar(select(Product).where(Product.trade_name==trade)) if trade else None
        price=cell(r,'price','price_egp','sale_price','السعر')
        try: price_f=float(Decimal(price)) if price else 0
        except (InvalidOperation,ValueError): price_f=0
        if not p:
            if not barcode:
                key='|'.join([trade,ar,generic,cell(r,'manufacturer','company','الشركة')]).encode('utf-8')
                barcode='EGD-'+hashlib.sha1(key).hexdigest()[:16].upper()
            p=Product(barcode=barcode,trade_name=trade or ar,active_ingredient=generic,strength=cell(r,'strength','التركيز'),dosage_form=cell(r,'dosage_form','form','route','الشكل الدوائي','طريقة الاستخدام'),manufacturer=cell(r,'manufacturer','company','الشركة'),sale_price=price_f); s.add(p); s.flush(); made+=1
        else:
            if trade:p.trade_name=trade
            if generic:p.active_ingredient=generic
            if price:p.sale_price=price_f
            for attr,names in [('strength',('strength','التركيز')),('dosage_form',('dosage_form','form','route','الشكل الدوائي','طريقة الاستخدام')),('manufacturer',('manufacturer','company','الشركة'))]:
                v=cell(r,*names)
                if v:setattr(p,attr,v)
            updated+=1
        d=s.get(DrugProfile,p.id) or DrugProfile(product_id=p.id); s.add(d)
        d.arabic_name=ar or d.arabic_name; d.generic_name=generic or d.generic_name; d.generic_name_ar=cell(r,'generic_name_ar','active_ar','المادة الفعالة عربي') or d.generic_name_ar
        d.category=cell(r,'category','drug_class','التصنيف') or d.category; d.indications=cell(r,'indications','دواعي الاستعمال') or d.indications; d.contraindications=cell(r,'contraindications','موانع الاستعمال') or d.contraindications; d.warnings=cell(r,'warnings','تحذيرات') or d.warnings; d.side_effects=cell(r,'side_effects','الآثار الجانبية') or d.side_effects; d.interactions=cell(r,'interactions','التفاعلات') or d.interactions; d.source=cell(r,'source','المصدر') or d.source; d.source_updated=cell(r,'source_updated','تاريخ المصدر') or d.source_updated
        level=norm(cell(r,'dispensing_level','level','حالة الصرف'))
        aliases={'otc':'otc','rx':'rx','prescription':'rx','restricted':'restricted','controlled':'controlled'}
        if level in aliases:
            pol=s.get(DispensingPolicy,p.id) or DispensingPolicy(product_id=p.id);s.add(pol);pol.level=aliases[level];pol.prescription_required=1 if pol.level!='otc' else 0
    s.add(Audit(action='DRUG_CATALOG_IMPORT',details=f'created={made} updated={updated} skipped={skipped}'));s.commit();return made,updated,skipped

@app.get('/drug-import',response_class=HTMLResponse)
def drug_import_page():
    return page('استيراد قاعدة الأدوية المصرية','''<div class="card"><h2>CSV / Excel</h2><p>يدعم: barcode/GTIN, trade_name, arabic_name, generic_name/active_ingredient, generic_name_ar, strength, dosage_form, manufacturer, category, price, warnings, side_effects, interactions, dispensing_level, source.</p><p><b>مهم:</b> بيانات التحذيرات والتفاعلات يجب أن تأتي من مصدر موثوق/مرخص ومذكور في عمود source.</p><form method="post" enctype="multipart/form-data"><input type="file" name="file" accept=".csv,.xlsx,.xlsm" required><button>استيراد / تحديث</button></form></div>''')

@app.post('/drug-import',response_class=HTMLResponse)
async def drug_import(file:UploadFile=File(...),s:Session=Depends(db)):
    data=await file.read()
    if len(data)>30*1024*1024: raise HTTPException(413,'الملف أكبر من 30MB')
    rows=rows_from_upload(file.filename or '',data); made,updated,skipped=import_catalog(s,rows)
    return page('نتيجة الاستيراد',f'<div class="card"><h2>تم الاستيراد</h2><p>جديد: <b>{made}</b> | تحديث: <b>{updated}</b> | متخطى: <b>{skipped}</b></p><a href="/drug-guide">فتح دليل الأدوية</a></div>')

@app.get('/smart-alternatives',response_class=HTMLResponse)
def smart_alternatives(q:str='',s:Session=Depends(db)):
    found=[]
    if q:
        like=f'%{q}%'; target=s.scalar(select(Product).where(or_(Product.trade_name.ilike(like),Product.barcode==q,Product.active_ingredient.ilike(like))).order_by(Product.id))
        if target:
            g=norm(target.active_ingredient); st=norm(target.strength)
            for p in s.scalars(select(Product).where(Product.id!=target.id)).all():
                if g and norm(p.active_ingredient)==g:
                    score=2+(1 if st and norm(p.strength)==st else 0)+(1 if norm(p.dosage_form)==norm(target.dosage_form) and p.dosage_form else 0)
                    found.append((score,p))
            found.sort(key=lambda x:(-x[0],x[1].sale_price,x[1].trade_name))
            tr=''.join(f'<tr><td>{escape(p.trade_name)}</td><td>{escape(p.active_ingredient)}</td><td>{escape(p.strength)}</td><td>{escape(p.dosage_form)}</td><td>{p.sale_price:.2f}</td><td>{score}/4</td></tr>' for score,p in found[:50])
            return page('البدائل والمثائل الذكية',f'<form><input name="q" value="{escape(q)}" placeholder="اسم / باركود / مادة فعالة"><button>بحث</button></form><div class="card"><b>الأصل:</b> {escape(target.trade_name)} — {escape(target.active_ingredient)} {escape(target.strength)}</div><table><tr><th>البديل</th><th>المادة</th><th>التركيز</th><th>الشكل</th><th>السعر</th><th>تطابق</th></tr>{tr}</table>')
    return page('البدائل والمثائل الذكية',f'<form><input name="q" value="{escape(q)}" placeholder="اسم / باركود / مادة فعالة"><button>بحث</button></form><p>يتم الاقتراح حسب تطابق المادة الفعالة، ثم التركيز والشكل الدوائي. القرار النهائي للصيدلي.</p>')

def ingredient_tokens(p):
    return {norm(x) for x in re.split(r'[/,+;]|\band\b',p.active_ingredient or '') if norm(x)}

def interaction_hits(s,products):
    rules=s.scalars(select(DrugInteraction)).all(); hits=[]
    for i,a in enumerate(products):
        for b in products[i+1:]:
            ta,tb=ingredient_tokens(a),ingredient_tokens(b)
            for r in rules:
                ra,rb=norm(r.ingredient_a),norm(r.ingredient_b)
                if (ra in ta and rb in tb) or (rb in ta and ra in tb): hits.append((a,b,r))
    return hits

@app.get('/interaction-checker',response_class=HTMLResponse)
def interaction_checker(ids:str='',s:Session=Depends(db)):
    products=[]
    for x in ids.split(','):
        if x.strip().isdigit():
            p=s.get(Product,int(x)); products.append(p) if p and p not in products else None
    opts=''.join(f'<option value="{p.id}">{escape(p.trade_name)} — {escape(p.active_ingredient)}</option>' for p in s.scalars(select(Product).order_by(Product.trade_name)).all())
    hits=interaction_hits(s,products); tr=''.join(f'<tr><td>{escape(a.trade_name)} + {escape(b.trade_name)}</td><td>{SEVERITY.get(r.severity,r.severity)}</td><td>{escape(r.note)}</td><td>{escape(r.source)}</td></tr>' for a,b,r in hits)
    selected='، '.join(escape(p.trade_name) for p in products) or 'لم يتم اختيار أدوية'
    return page('فحص التفاعلات الدوائية',f'''<div class="card"><p>اختر دواء ثم أضف رقمه إلى القائمة (مثال: 1,2,3). الفحص يعتمد فقط على قواعد موثقة تمت إضافتها للنظام ولا يغني عن المرجع الدوائي.</p><select id="drug">{opts}</select><button type="button" onclick="let u=new URL(location);let a=(u.searchParams.get('ids')||'').split(',').filter(Boolean);let v=document.getElementById('drug').value;if(!a.includes(v))a.push(v);u.searchParams.set('ids',a.join(','));location=u">إضافة للفحص</button><a href="/interaction-checker">مسح</a><p><b>المحدد:</b> {selected}</p></div><table><tr><th>الأدوية</th><th>الشدة</th><th>الملاحظة</th><th>المصدر</th></tr>{tr or '<tr><td colspan="4">لا توجد تفاعلات مسجلة بين الاختيارات في قاعدة النظام.</td></tr>'}</table>''')

@app.get('/drug-interactions',response_class=HTMLResponse)
def interactions_admin(s:Session=Depends(db)):
    rows=s.scalars(select(DrugInteraction).order_by(DrugInteraction.id.desc()).limit(200)).all();tr=''.join(f'<tr><td>{escape(x.ingredient_a)}</td><td>{escape(x.ingredient_b)}</td><td>{SEVERITY.get(x.severity,x.severity)}</td><td>{escape(x.note)}</td><td>{escape(x.source)}</td></tr>' for x in rows)
    return page('قواعد التفاعلات الدوائية',f'''<div class="card"><form method="post"><input name="ingredient_a" placeholder="Active ingredient A" required><input name="ingredient_b" placeholder="Active ingredient B" required><select name="severity"><option value="minor">Minor</option><option value="moderate">Moderate</option><option value="major">Major</option><option value="contraindicated">Contraindicated</option></select><input name="note" placeholder="التأثير/الإجراء"><input name="source" placeholder="المصدر الموثوق" required><button>إضافة قاعدة</button></form></div><table><tr><th>A</th><th>B</th><th>الشدة</th><th>الملاحظة</th><th>المصدر</th></tr>{tr}</table>''')

@app.post('/drug-interactions')
def interaction_add(ingredient_a:str=Form(...),ingredient_b:str=Form(...),severity:str=Form(...),note:str=Form(''),source:str=Form(...),s:Session=Depends(db)):
    if severity not in SEVERITY or norm(ingredient_a)==norm(ingredient_b) or not source.strip(): raise HTTPException(400,'بيانات غير صحيحة')
    s.add(DrugInteraction(ingredient_a=ingredient_a.strip(),ingredient_b=ingredient_b.strip(),severity=severity,note=note.strip(),source=source.strip()));s.add(Audit(action='INTERACTION_RULE_ADD',details=f'{ingredient_a} + {ingredient_b}'));s.commit();return RedirectResponse('/drug-interactions',303)
