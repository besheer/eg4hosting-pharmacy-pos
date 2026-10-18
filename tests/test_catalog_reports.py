import io, csv, json, datetime as dt
from test_pos import client, engine
from sqlalchemy.orm import Session
from sqlalchemy import select
from app.main import Product, Batch, Sale, SaleItem, SalesReturn, User, Role, password_hash
from app.pos_v05 import SaleTotals, ItemAmounts, Payment
from app.catalog_reports import ImportJob, workbook_bytes
from fastapi.testclient import TestClient
from app.main import app
from openpyxl import load_workbook


def upload(content,name='products.csv',mode='add'):
    r=client.post('/catalog/preview',files={'file':(name,content)},data={'mode':mode},follow_redirects=False)
    assert r.status_code==303,r.text
    return r.headers['location'].split('/')[-1]
def confirm(token):return client.post('/catalog/confirm/'+token,follow_redirects=False)
def job(token):
    with Session(engine) as s:return json.loads(s.get(ImportJob,token).payload)


def test_import_preview_confirm_replay_and_excel_export(tmp_path,monkeypatch):
    monkeypatch.chdir(tmp_path)
    raw='barcode,trade_name,sale_price,reorder_level\n00000042,دواء تجربة,12.50,4\n'.encode('utf-8-sig')
    token=upload(raw)
    with Session(engine) as s:assert s.scalar(select(Product).where(Product.barcode=='00000042')) is None
    assert not job(token)['errors']
    assert client.get('/catalog/preview/'+token).status_code==200
    assert confirm(token).status_code==303
    assert confirm(token).status_code==303
    assert len(list((tmp_path/'backups').glob('*.db')))==1
    with Session(engine) as s:
        p=s.scalar(select(Product).where(Product.barcode=='00000042'));assert p.trade_name=='دواء تجربة' and p.sale_price==12.5
    result=client.get('/catalog/export?fmt=xlsx&q=00000042');assert result.status_code==200
    wb=load_workbook(io.BytesIO(result.content));assert wb.active['A2'].value=='00000042';assert wb.active['A2'].data_type=='s'
    assert wb.active['G2'].value==12.5
    token=upload(result.content,'roundtrip.xlsx','upsert');assert not job(token)['errors'];assert confirm(token).status_code==303
    result=client.get('/catalog/export?fmt=csv&q=00000042');assert result.content.startswith(b'\xef\xbb\xbf');assert 'دواء تجربة' in result.content.decode('utf-8-sig')


def test_import_validation_stale_and_atomic(tmp_path,monkeypatch):
    monkeypatch.chdir(tmp_path)
    token=upload(b'barcode,trade_name,sale_price\n0099,valid,5\n0099,duplicate,9\n0010,bad,-1\n')
    assert len(job(token)['errors'])==2
    assert confirm(token).status_code==400
    assert client.get('/catalog/errors/'+token).status_code==200
    with Session(engine) as s:assert s.scalar(select(Product).where(Product.barcode=='0099')) is None
    token=upload(b'barcode,trade_name,sale_price\n00000042,changed,30\n','update.csv','upsert')
    with Session(engine) as s:
        p=s.scalar(select(Product).where(Product.barcode=='00000042'));p.sale_price=99;s.commit()
    assert confirm(token).status_code==409
    with Session(engine) as s:assert s.scalar(select(Product).where(Product.barcode=='00000042')).sale_price==99
    numeric=workbook_bytes(['barcode','trade_name','sale_price'],[[123,'bad numeric barcode',1]])
    token=upload(numeric,'numeric.xlsx');assert job(token)['errors']
    assert client.post('/catalog/preview',files={'file':('bad.xls',b'x')}).status_code==400


def test_xlsx_formula_rejected_and_export_is_literal():
    payload=workbook_bytes(['barcode','trade_name','sale_price'],[['0044','=1+1',1]])
    wb=load_workbook(io.BytesIO(payload));assert wb.active['B2'].data_type=='s'
    wb.active['C2']='=1+1';buf=io.BytesIO();wb.save(buf)
    assert client.post('/catalog/preview',files={'file':('formula.xlsx',buf.getvalue())}).status_code==400


def test_reports_event_dates_cost_tax_and_branch():
    with Session(engine) as s:
        p=Product(barcode='report-test',trade_name='Report drug',sale_price=100);s.add(p);s.flush()
        b=Batch(product_id=p.id,branch_id=1,batch_no='report',expiry=dt.date(2031,1,1),qty=2,cost=999);s.add(b);s.flush()
        sale=Sale(branch_id=1,total=198,created_at=dt.datetime(2030,1,1,12));s.add(sale);s.flush()
        i=SaleItem(sale_id=sale.id,product_id=p.id,batch_id=b.id,qty=2,unit_price=100);s.add(i);s.flush()
        s.add(ItemAmounts(item_id=i.id,net=18000,tax=1800,cost=12000))
        s.add(SaleTotals(sale_id=sale.id,request_key='report-fixture-key',request_hash='x',user_id=1,subtotal=20000,discount=2000,tax=1800,total=19800,tax_pct='10',snapshot='{}'))
        s.add(Payment(sale_id=sale.id,method='cash',amount=19800))
        s.add(SalesReturn(sale_id=sale.id,sale_item_id=i.id,branch_id=1,qty=1,amount=99,created_at=dt.datetime(2030,2,1,12)));s.commit()
    def rows(start,end,kind='sales'):
        r=client.get('/reports',params={'start':start,'end':end,'kind':kind,'q':'report-test','fmt':'csv'})
        assert r.status_code==200,r.text
        return list(csv.reader(io.StringIO(r.content.decode('utf-8-sig'))))
    jan=rows('2030-01-01','2030-01-31');assert len(jan)==2
    assert list(map(float,jan[1][8:]))==[180,18,198,120,60]
    feb=rows('2030-02-01','2030-02-28');assert len(feb)==2
    assert list(map(float,feb[1][8:]))==[-90,-9,-99,-60,-30]
    both=rows('2030-01-01','2030-02-28','profit');assert list(map(float,both[1][2:]))==[1,90,60,30]
    assert client.get('/reports?start=2030-02-01&end=2030-01-01').status_code==400
    for kind in ['sales','profit','payments','inventory','low','expiry','ledger','purchases']:
        assert client.get('/reports',params={'kind':kind}).status_code==200
        r=client.get('/reports',params={'kind':kind,'fmt':'xlsx'});assert r.status_code==200
        assert load_workbook(io.BytesIO(r.content)).active.max_row>=1
    with Session(engine) as s:
        role=s.scalar(select(Role).where(Role.name=='Accountant'));s.add(User(username='report-user',password_hash=password_hash('pw'),role_id=role.id,branch_id=2));s.commit()
    other=TestClient(app);other.post('/login',data={'username':'report-user','password':'pw'})
    assert other.get('/reports?branch_id=1').status_code==403
    assert other.get('/catalog').status_code==403
    assert other.get('/catalog/export?fmt=csv').status_code==403
    assert other.get('/reports?kind=sales&start=2030-01-01&end=2030-02-28&fmt=csv').content.decode('utf-8-sig').count('\n')==1
